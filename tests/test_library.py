"""Tests for reading many files at once and reconciling them.

Runs under pytest or standalone:  python -m tests.test_library

The reconciliation rules are the whole point, and each of them can be wrong in a
way that produces a plausible number nobody questions — doubled stock, doubled
revenue, a snapshot added to itself. Those are what these test.
"""

from __future__ import annotations

import os
import shutil
import sys
from pathlib import Path

os.environ["VYUHA_LLM"] = "offline"

from fastapi.testclient import TestClient                  # noqa: E402
from openpyxl import Workbook                              # noqa: E402

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

from vyuha_platform import (app as app_mod, auth, config, intake,  # noqa: E402
                            library, store)

client = TestClient(app_mod.app, follow_redirects=True)
WORK = REPO / "out" / "lib-test"
SETTINGS = config.load()
_SLUGS: list[str] = []

ACCOUNT = auth.create(f"lib-{auth.secrets.token_hex(4)}@vyuha.test",
                      "Library Test Operator", "test-password-1")


def _login() -> None:
    client.post("/login", data={"email": ACCOUNT.email, "password": "test-password-1"})


def _as_operator() -> None:
    a = auth.get(ACCOUNT.id)
    a.install, a.org_name, a.tenant_slug = "operator", "", ""
    auth.update(a)


_login()
_as_operator()


def _dir(name: str) -> Path:
    d = WORK / name
    shutil.rmtree(d, ignore_errors=True)
    d.mkdir(parents=True, exist_ok=True)
    return d


def _sales(path: Path, rows: list[tuple], headers=None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(headers or ["Date", "Invoice No", "Party", "Item", "Qty", "Rate", "Amount"])
    for r in rows:
        ws.append(list(r))
    wb.save(path)
    return path


def _stock(path: Path, rows: list[tuple]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(["Item Code", "Description", "Closing Qty", "Reorder Level", "Rate"])
    for r in rows:
        ws.append(list(r))
    wb.save(path)
    return path


def _read(folder: Path) -> library.Batch:
    files, _notes = library.scan(folder)
    return library.batch(files, WORK / "wk", SETTINGS, label="test")


# ================================================== the reconciliation rules

def test_sales_from_different_files_add_up():
    d = _dir("accumulate")
    _sales(d / "jan.xlsx", [("2026-01-05", "A-1", "Ramu", "Urea", 10, 320, 3200)])
    _sales(d / "feb.xlsx", [("2026-02-05", "A-2", "Ramu", "Urea", 5, 320, 1600)])

    b = _read(d)
    assert b.read == 2
    assert (b.insights.sales.get("revenue") or 0) == 4800


def test_the_same_bill_in_two_files_is_counted_once():
    """Client exports overlap constantly — a 90-day report sent twice shares a
    month, and stacking those doubles a third of the revenue."""
    d = _dir("dedupe")
    shared = ("2026-01-05", "A-1", "Ramu", "Urea", 10, 320, 3200)
    _sales(d / "q1.xlsx", [shared, ("2026-02-01", "A-2", "Ramu", "Urea", 5, 320, 1600)])
    _sales(d / "q1-again.xlsx", [shared])

    b = _read(d)
    assert b.duplicates_dropped == 1, b.duplicates_dropped
    assert (b.insights.sales.get("revenue") or 0) == 4800, "the shared bill was doubled"


def test_a_row_with_no_invoice_number_is_still_deduped():
    d = _dir("dedupe-noinv")
    headers = ["Date", "Party", "Item", "Qty", "Rate", "Amount"]
    row = ("2026-01-05", "Ramu", "Urea", 10, 320, 3200)
    _sales(d / "a.xlsx", [row], headers)
    _sales(d / "b.xlsx", [row], headers)

    b = _read(d)
    assert b.duplicates_dropped == 1
    assert (b.insights.sales.get("revenue") or 0) == 3200


def test_stock_is_a_snapshot_and_the_newest_file_wins():
    """The most expensive mistake this module could make: a March count and an
    August count are not two sets of shelves."""
    d = _dir("stock-snapshot")
    old = _stock(d / "march.xlsx", [("URE-50", "Urea 50kg", 100, 20, 320)])
    new = _stock(d / "august.xlsx", [("URE-50", "Urea 50kg", 40, 20, 320)])
    os.utime(old, (1_600_000_000, 1_600_000_000))
    os.utime(new, (1_700_000_000, 1_700_000_000))

    b = _read(d)
    stock = next(t for t in b.tables if t.kind == "stock")
    qty = float(stock.frame["stock_qty"].sum())
    assert qty == 40, f"stock was added instead of replaced: {qty}"
    assert b.stock_files_superseded == 1


def test_an_item_only_in_an_older_stock_file_is_not_lost():
    """Newest wins per SKU — it does not delete lines the new file omits."""
    d = _dir("stock-carry")
    old = _stock(d / "march.xlsx", [("URE-50", "Urea", 100, 20, 320),
                                    ("DAP-50", "DAP", 50, 10, 1420)])
    new = _stock(d / "august.xlsx", [("URE-50", "Urea", 40, 20, 320)])
    os.utime(old, (1_600_000_000, 1_600_000_000))
    os.utime(new, (1_700_000_000, 1_700_000_000))

    b = _read(d)
    stock = next(t for t in b.tables if t.kind == "stock")
    skus = set(stock.frame["sku"].astype(str))
    assert skus == {"URE-50", "DAP-50"}, skus
    assert float(stock.frame["stock_qty"].sum()) == 90      # 40 new + 50 carried


def test_every_file_gets_a_verdict_whether_it_read_or_not():
    """A hundred files in, a hundred lines out. Silence about a dropped file is
    worse than a failure, because nobody notices."""
    d = _dir("verdicts")
    _sales(d / "good.xlsx", [("2026-01-05", "A-1", "Ramu", "Urea", 10, 320, 3200)])
    (d / "bad.xlsx").write_bytes(b"%PDF-1.4\nnot a workbook")
    (d / "empty.csv").write_text("just,some,words\nno,numbers,here\n", encoding="utf-8")

    b = _read(d)
    assert len(b.files) == 3
    assert b.read >= 1 and b.rejected >= 1
    for f in b.files:
        assert f.ok or f.error, f"{f.name} had no verdict"
    bad = next(f for f in b.files if f.name == "bad.xlsx")
    assert "PDF" in bad.error, bad.error


def test_one_unreadable_file_does_not_stop_the_rest():
    d = _dir("resilient")
    for i in range(4):
        _sales(d / f"ok{i}.xlsx",
               [("2026-01-05", f"A-{i}", "Ramu", "Urea", 10, 320, 3200)])
    (d / "wreck.xlsx").write_bytes(b"\x00" * 200)

    b = _read(d)
    assert b.read == 4
    assert (b.insights.sales.get("revenue") or 0) == 12800


def test_scanning_a_folder_finds_files_and_says_what_it_left():
    d = _dir("scan")
    _sales(d / "a.xlsx", [("2026-01-05", "A-1", "R", "Urea", 1, 320, 320)])
    (d / "notes.docx").write_text("not readable", encoding="utf-8")
    (d / "~$a.xlsx").write_text("excel lock file", encoding="utf-8")

    found, notes = library.scan(d)
    assert [p.name for p in found] == ["a.xlsx"]
    assert any("other types" in n for n in notes)


def test_scanning_a_missing_folder_says_so_rather_than_failing():
    found, notes = library.scan(Path("no/such/folder/anywhere"))
    assert found == []
    assert any("no folder" in n.lower() for n in notes)


# ============================================================== the routes

def _shop(name: str) -> str:
    client.post("/onboard", data={"name": name, "phone": "9876543210",
                                  "data_mode": "upload"})
    slug = next(c.slug for c in store.load_clients(ACCOUNT.id)
                if c.slug.startswith(store.slugify(name)))
    _SLUGS.append(slug)
    return slug


def test_many_files_upload_in_one_go():
    """Selecting ninety files and getting the ninetieth back was the bug."""
    slug = _shop("Multi Upload Co")
    d = _dir("upload")
    paths = [_sales(d / f"m{i}.xlsx",
                    [(f"2026-0{i+1}-05", f"A-{i}", "Ramu", "Urea", 10, 320, 3200)])
             for i in range(3)]

    files = [("file", (p.name, p.open("rb"), "application/octet-stream"))
             for p in paths]
    resp = client.post(f"/c/{slug}/upload", files=files)
    assert resp.status_code == 200

    run = store.get_client(slug, ACCOUNT.id).latest
    assert run.status == "ok"
    assert run.revenue == 9600, f"only one file was read: {run.revenue}"
    assert "3 files" in run.filename


def test_a_folder_can_be_read_where_it_sits():
    slug = _shop("Folder Read Co")
    d = _dir("folderroute")
    _sales(d / "a.xlsx", [("2026-01-05", "A-1", "Ramu", "Urea", 10, 320, 3200)])
    _sales(d / "b.xlsx", [("2026-02-05", "A-2", "Ramu", "Urea", 10, 320, 3200)])

    resp = client.post(f"/c/{slug}/folder", data={"path": str(d), "recursive": "1"})
    assert resp.status_code == 200
    run = store.get_client(slug, ACCOUNT.id).latest
    assert run.revenue == 6400

    # The folder is read, never moved.
    assert (d / "a.xlsx").exists()


def test_a_folder_that_is_not_there_is_reported_not_crashed():
    slug = _shop("Bad Folder Co")
    resp = client.post(f"/c/{slug}/folder", data={"path": "Z:/nope/nothing"})
    assert resp.status_code == 200
    assert "no folder" in resp.text.lower() or "nothing readable" in resp.text.lower()


def test_uploading_nothing_says_so():
    slug = _shop("Empty Upload Co")
    resp = client.post(f"/c/{slug}/upload", files=[])
    assert resp.status_code == 200


# ====================================================== the chat, catalogue-free

def test_a_whatsapp_thread_reads_without_a_catalogue():
    """The bug on screen: a business that sends files has an empty Book, so
    nothing matched, so a successfully parsed thread reported as unreadable."""
    text = (REPO / "demo" / "samples" / "08-whatsapp-orders.txt").read_text(
        encoding="utf-8")
    extract = intake.parse_chat(text, [], SETTINGS)
    assert extract.ok, extract.error
    assert len(extract.orders) >= 5, [d.item for d in extract.orders]
    # Names lifted from a sentence are guesses and must say so.
    assert all(d.confidence == "low" for d in extract.orders)
    assert all(d.item.strip() for d in extract.orders)


def test_a_catalogue_still_wins_when_there_is_one():
    text = (REPO / "demo" / "samples" / "08-whatsapp-orders.txt").read_text(
        encoding="utf-8")
    extract = intake.parse_chat(text, ["Urea 50kg", "Gypsum 5kg"], SETTINGS)
    named = [d for d in extract.orders if d.item in {"Urea 50kg", "Gypsum 5kg"}]
    assert named, [d.item for d in extract.orders]
    assert all(d.confidence == "medium" for d in named)


def test_a_thread_with_no_orders_is_reported_not_handed_on_empty():
    """It used to write a header-only CSV, and the engine then reported a parse
    failure about a file that had parsed perfectly."""
    chat = "\n".join([
        "01/09/2026, 9:03 am - Ramu Stores: Namaskara sir",
        "01/09/2026, 9:04 am - You: Namaskara",
        "01/09/2026, 9:05 am - Ramu Stores: how is the family",
        "01/09/2026, 9:06 am - You: all well sir",
    ])
    extract = intake.parse_chat(chat, [], SETTINGS)
    assert not extract.ok
    assert "no orders" in extract.error.lower() or "nothing" in extract.error.lower()


# ================================================ uploads reach every screen

def _upload_to(slug: str, names: list[str]):
    files = [("file", (n, (REPO / "demo" / "samples" / n).open("rb"),
                       "application/octet-stream")) for n in names]
    return client.post(f"/c/{slug}/upload", files=files)


def test_an_uploaded_file_reaches_the_stock_screen():
    """The bug: files parsed perfectly and every screen showed zero, because
    the console reads the Book and only typed entries ever wrote one."""
    from vyuha_platform import books
    slug = _shop("End To End Co")
    _upload_to(slug, ["02-filthy-multisheet.xlsx"])

    book = books.load(slug)
    assert book.items, "nothing landed in the book"
    assert book.sales, "no sales landed in the book"

    stock = client.get(f"/c/{slug}/stock").text
    assert 'class="sk ' in stock, "the shelf is empty"
    assert book.items[0].name.split()[0] in stock


def test_an_upload_gives_every_item_a_price():
    """A stock statement rarely carries one, and an item at zero makes stock
    value, margin and the cost of a stockout all come out as zero."""
    from vyuha_platform import books
    slug = _shop("Priced Co")
    _upload_to(slug, ["01-clean-sales.csv", "02-filthy-multisheet.xlsx"])

    book = books.load(slug)
    priced = [i for i in book.items if i.rate > 0]
    assert len(priced) == len(book.items), [i.name for i in book.items if not i.rate]
    assert book.stock_value > 0


def test_one_customer_spelled_four_ways_is_one_customer():
    """The cleaner works this out and leaves it in party_key; writing the raw
    spelling threw it away and split a 34% customer into 19% and 15% — under
    the threshold, so a real concentration risk went unflagged."""
    from vyuha_platform import books, finance, money as money_mod
    slug = _shop("One Name Co")
    _upload_to(slug, ["02-filthy-multisheet.xlsx"])

    book = books.load(slug)
    spellings = {s.party.strip().lower().rstrip(".") for s in book.sales
                 if "ramu" in s.party.lower()}
    assert len(spellings) == 1, spellings

    conc = finance.concentration(book, money_mod.load(slug))
    assert conc["customers"][0]["share"] > 0.25


def test_today_speaks_about_uploaded_data_in_rupees():
    """Not "4 SKU(s) at or below reorder level" — a sentence with money in it."""
    from vyuha_platform import books, invoice, money as money_mod, people, today
    slug = _shop("Findings Co")
    _upload_to(slug, ["01-clean-sales.csv", "02-filthy-multisheet.xlsx"])

    c = store.get_client(slug, ACCOUNT.id)
    found = today.findings(c, books.load(slug), money_mod.load(slug),
                           people.load(slug), invoice.load_all(slug))
    assert found, "nothing surfaced from an uploaded file"
    stock = next((f for f in found if "stock" in f.tags), None)
    assert stock is not None
    assert "₹0" not in stock.detail, "cost of a stockout came out as zero"


def test_uploading_again_replaces_rather_than_doubles():
    """Files are the source of truth for this kind of business, so a re-upload
    is a correction, not an addition."""
    from vyuha_platform import books
    slug = _shop("Reupload Co")
    _upload_to(slug, ["01-clean-sales.csv"])
    first = books.load(slug).earned
    _upload_to(slug, ["01-clean-sales.csv"])
    again = books.load(slug).earned
    assert again == first, f"{first} became {again}"


def test_an_upload_never_deletes_typed_entries():
    """Losing hand-entered sales to a stray file would be unforgivable."""
    from vyuha_platform import books
    client.post("/onboard", data={"name": "Typed Co", "phone": "9876543210",
                                  "data_mode": "books"})
    slug = next(c.slug for c in store.load_clients(ACCOUNT.id)
                if c.slug.startswith("typed-co"))
    _SLUGS.append(slug)
    client.post(f"/c/{slug}/book/item", data={
        "name": "Hand Typed Item", "category": "Other", "unit": "piece",
        "rate": "500", "cost": "300", "stock_qty": "10", "reorder_level": "2"})

    _upload_to(slug, ["01-clean-sales.csv"])
    names = {i.name for i in books.load(slug).items}
    assert "Hand Typed Item" in names, names


def _cleanup() -> None:
    _login()
    _as_operator()
    for slug in set(_SLUGS):
        shutil.rmtree(store.UPLOADS / slug, ignore_errors=True)
        shutil.rmtree(store.DASHBOARDS / slug, ignore_errors=True)
        store.delete_client(slug, ACCOUNT.id)
    shutil.rmtree(WORK, ignore_errors=True)


def main() -> int:
    tests = [(n, f) for n, f in sorted(globals().items())
             if n.startswith("test_") and callable(f)]
    passed = failed = 0
    try:
        for name, fn in tests:
            try:
                fn()
                print(f"ok   {name}")
                passed += 1
            except Exception as exc:                       # noqa: BLE001
                print(f"FAIL {name}: {type(exc).__name__}: {exc}")
                failed += 1
    finally:
        _cleanup()
    print(f"\n{passed} passed, {failed} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
