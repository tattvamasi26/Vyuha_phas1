"""Tests for the new site (/app, /studio) and the Phase 1 data fixes.

Runs under pytest or standalone:  python -m tests.test_web

Like the other platform suites this drives the real routes against the real
``vyuha_data/``, and removes every account and business it creates in
``_cleanup()``. Assertions look at ``data-testid`` hooks and at the data, not at
wording, so rewriting a sentence on a page never breaks a test.
"""

from __future__ import annotations

import os
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

os.environ["VYUHA_LLM"] = "offline"          # before the platform imports llm

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO))

import pandas as pd                                                        # noqa: E402
from fastapi.testclient import TestClient                                  # noqa: E402

from vyuha import clean, pipeline                                          # noqa: E402
from vyuha_platform import (app as app_mod, auth, books, invoice, library,  # noqa: E402
                            onboarding, store)
from vyuha_platform.web import nav                                         # noqa: E402

client = TestClient(app_mod.app, follow_redirects=True)
_SLUGS: list[str] = []
PASSWORD = "test-password-1"

ACCOUNT = auth.create(f"web-{auth.secrets.token_hex(4)}@vyuha.test", "Web Test Operator",
                      PASSWORD)
OTHER = auth.create(f"web-other-{auth.secrets.token_hex(4)}@vyuha.test", "Someone Else",
                    PASSWORD)


def _install(acct, kind: str = "operator", slug: str = "") -> None:
    a = auth.get(acct.id)
    a.install, a.tenant_slug = kind, slug
    auth.update(a)


def _login(acct=ACCOUNT) -> None:
    client.cookies.clear()
    client.post("/login", data={"email": acct.email, "password": PASSWORD})


_install(ACCOUNT)
_install(OTHER)
_login()


def _onboard(name: str, **extra) -> str:
    resp = client.post("/studio/new", data={"name": name, "contact": "Ravi Kumar",
                                            "phone": "98450 12345", **extra})
    assert resp.status_code == 200, resp.status_code
    slug = next(c.slug for c in store.load_clients(ACCOUNT.id) if c.name == name)
    _SLUGS.append(slug)
    return slug


def _stock(slug: str) -> None:
    """A little shop: two items, one running low, a cash sale and a credit sale."""
    books.add_item(slug, "Urea 50kg", "Fertiliser", "bag", 320, 268, 40, 20)
    books.add_item(slug, "DAP 50kg", "Fertiliser", "bag", 1420, 1250, 3, 15)
    book = books.load(slug)
    urea = next(i for i in book.items if i.name == "Urea 50kg")
    books.record_sale(slug, urea.sku, "Ramu Stores", 5, 320)
    books.record_sale(slug, urea.sku, "Basavaraj Agri", 2, 320, paid=False,
                      due_date="2026-01-01")


# ================================================== the Phase 1 data fixes

def test_iso_dates_are_read_as_iso_and_the_rest_day_first():
    """2026-08-12 is 12 August. Read day-first it became 8 December."""
    raw = pd.Series(["2026-08-12", "2026-09-01", "12/08/2026", "05-11-2025", ""])
    out = clean._to_dates(raw)
    got = [d.strftime("%Y-%m-%d") if not pd.isna(d) else "" for d in out]
    assert got == ["2026-08-12", "2026-09-01", "2026-08-12", "2025-11-05", ""], got


def test_a_typed_book_goes_to_the_engine_as_real_dates():
    """Dates leave the book as date cells, so nothing downstream guesses the month."""
    from openpyxl import load_workbook
    book = books.Book(slug="x", sales=[books.Sale(
        id="B-1", date="2026-08-12", party="Ramu", sku="U", item="Urea", qty=1,
        rate=10, amount=10)])
    with tempfile.TemporaryDirectory() as tmp:
        path = books.to_workbook(book, Path(tmp) / "b.xlsx")
        cell = load_workbook(path)["Sales Register"]["A2"].value
    assert isinstance(cell, datetime) and cell.date() == date(2026, 8, 12), cell


def test_an_invoice_is_for_one_customer():
    slug = _onboard("Mixed Bill Traders")
    _stock(slug)
    book = books.load(slug)
    ids = [s.id for s in book.sales]
    assert len({s.party for s in book.sales}) == 2
    try:
        invoice.from_sales(store.get_client(slug, ACCOUNT.id), book, ids)
    except ValueError as exc:
        assert "one customer" in str(exc)
    else:
        raise AssertionError("an invoice covering two customers was raised")


def test_sending_the_same_file_twice_does_not_double_the_sales():
    slug = _onboard("Twice Sent Traders")
    with tempfile.TemporaryDirectory() as tmp:
        csv = Path(tmp) / "sales.csv"
        csv.write_text("Date,Invoice No,Party Name,Item,Qty,Rate,Amount\n"
                       "2026-08-12,INV-1,Ramu Stores,Urea 50kg,10,320,3200\n"
                       "2026-08-13,INV-2,Basavaraj,DAP 50kg,5,1420,7100\n", encoding="utf-8")
        tables = pipeline.run(csv).tables
    library.materialise(tables, slug, replace=False)
    first = len(books.load(slug).sales)
    counts = library.materialise(tables, slug, replace=False)
    assert first == 2, first
    assert len(books.load(slug).sales) == 2, "the second send doubled the sales"
    assert counts["duplicates"] == 2, counts
    assert sorted(s.date for s in books.load(slug).sales) == ["2026-08-12", "2026-08-13"]


# ======================================================= the Studio

def test_a_new_business_lands_on_the_data_map():
    resp = client.post("/studio/new", data={"name": "Studio Fresh Traders", "contact": "Anita",
                                            "phone": "9845000999"})
    slug = next(c.slug for c in store.load_clients(ACCOUNT.id) if c.name == "Studio Fresh Traders")
    _SLUGS.append(slug)
    assert str(resp.url).split("?")[0].endswith(f"/studio/{slug}/datamap"), resp.url
    assert 'data-testid="studio-stage-datamap"' in resp.text
    record = onboarding.load(slug)
    assert record.cutover == date.today().replace(day=1).isoformat()
    assert store.get_client(slug, ACCOUNT.id).data_mode == "books"


def test_the_studio_lists_every_business_with_its_stage():
    slug = _onboard("Listed Traders")
    resp = client.get("/studio")
    assert resp.status_code == 200
    assert 'data-testid="studio"' in resp.text
    assert f"/studio/{slug}/" in resp.text


def test_the_profile_saves_and_completes_its_stage():
    slug = _onboard("Profile Traders")
    resp = client.post(f"/studio/{slug}/profile", data={
        "name": "Profile Traders Pvt", "contact": "Meena", "phone": "98450 11111",
        "trade": "retail", "gstin": "29abcde1234f1z5", "state": "KA",
        "address": "MG Road", "fy_start_month": "4", "advance": "1"})
    c = store.get_client(slug, ACCOUNT.id)
    assert (c.name, c.contact, c.phone, c.gstin, c.state) == \
        ("Profile Traders Pvt", "Meena", "919845011111", "29ABCDE1234F1Z5", "KA")
    assert str(resp.url).split("?")[0].endswith("/datamap"), "Save and continue did not advance"
    states = onboarding.status(onboarding.load(slug), c, books.load(slug), None)
    assert states["profile"] == "done"


def test_an_unknown_state_is_not_stored():
    slug = _onboard("Nowhere Traders")
    client.post(f"/studio/{slug}/profile", data={"name": "Nowhere Traders", "state": "ZZ"})
    assert store.get_client(slug, ACCOUNT.id).state == ""


def test_the_data_map_saves_answers_and_decides_how_the_business_is_read():
    slug = _onboard("Mapped Traders")
    data = {f"src_{k}": "paper" for k in onboarding.DOMAIN_KEYS}
    data.update({"src_sales": "excel", "since_sales": "2025-04", "keeper_sales": "son",
                 "cadence_sales": "daily", "src_purchases": "software",
                 "software_purchases": "Busy", "cutover": "2026-09-01",
                 "history_months": "12"})
    client.post(f"/studio/{slug}/datamap", data=data)
    record = onboarding.load(slug)
    assert record.answered == len(onboarding.DOMAIN_KEYS)
    assert record.answer("sales").since == "2025-04"
    assert record.answer("purchases").source_label == "Busy"
    assert record.cutover == "2026-09-01"
    # Sales kept in Excel and nothing recorded yet: this business sends files.
    assert store.get_client(slug, ACCOUNT.id).data_mode == "upload"
    plan = {row["key"]: row for row in onboarding.plan(record)}
    assert "Tally" not in plan["sales"]["history"] and plan["sales"]["source"] == "Excel / Sheets"


def test_a_bad_source_value_is_ignored():
    slug = _onboard("Odd Source Traders")
    client.post(f"/studio/{slug}/datamap", data={"src_sales": "carrier pigeon"})
    assert onboarding.load(slug).answer("sales").source == ""


def test_starter_items_are_added_and_finish_the_masters_stage():
    slug = _onboard("Starter Traders", trade="nursery")
    client.post(f"/studio/{slug}/masters/catalogue",
                data={"pick": ["0", "1"], "rate_0": "500"})
    book = books.load(slug)
    assert len(book.items) == 2, [i.name for i in book.items]
    assert book.items[0].rate == 500
    c = store.get_client(slug, ACCOUNT.id)
    assert onboarding.status(onboarding.load(slug), c, book, None)["masters"] == "done"


def test_the_item_grid_adds_typed_items_and_skips_blank_rows():
    slug = _onboard("Grid Traders")
    client.post(f"/studio/{slug}/masters/items", data={
        "name": ["Hammer", "", "Screws 100pc"], "category": ["Tools", "", "Fasteners"],
        "unit": ["piece", "", "packet"], "rate": ["340", "", "130"],
        "cost": ["240", "", "88"], "reorder": ["5", "", "10"]})
    names = sorted(i.name for i in books.load(slug).items)
    assert names == ["Hammer", "Screws 100pc"], names


def test_branches_and_staff_can_be_added_from_the_studio():
    slug = _onboard("Crew Traders")
    client.post(f"/studio/{slug}/masters/branch", data={"name": "Hubballi", "place": "Market Yard"})
    client.post(f"/studio/{slug}/masters/staff", data={"name": "Manju", "role": "Manager",
                                                        "phone": "98000 00001"})
    from vyuha_platform import people
    org = people.load(slug)
    assert [b.name for b in org.branches] == ["Hubballi"]
    assert [(s.name, s.role, s.phone) for s in org.staff] == [("Manju", "Manager", "919800000001")]


def test_every_stage_page_renders():
    slug = _onboard("Stage Walk Traders")
    for stage in onboarding.STAGES:
        resp = client.get(f"/studio/{slug}/{stage.key}")
        assert resp.status_code == 200, (stage.key, resp.status_code)
        assert f'data-testid="studio-stage-{stage.key}"' in resp.text, stage.key


# ======================================================= the site

def test_home_shows_the_ranked_list_and_the_figures():
    slug = _onboard("Home Traders")
    _stock(slug)
    resp = client.get(f"/app/{slug}")
    assert resp.status_code == 200
    for hook in ("home", "needs-you", "kpis", "kpi-sales", "kpi-stock", "finding"):
        assert f'data-testid="{hook}"' in resp.text, hook


def test_every_section_page_renders():
    slug = _onboard("Section Walk Traders")
    for sec in nav.SECTIONS[1:]:
        resp = client.get(f"/app/{slug}/{sec.key}")
        assert resp.status_code == 200, (sec.key, resp.status_code)
        assert f'data-testid="section-{sec.key}"' in resp.text, sec.key
        first = sec.pages[0].key
        assert str(resp.url).endswith(f"/{sec.key}/{first}"), resp.url


def test_an_unknown_section_goes_home():
    slug = _onboard("Lost Traders")
    resp = client.get(f"/app/{slug}/nowhere")
    assert str(resp.url).endswith(f"/app/{slug}"), resp.url


def test_the_assistant_answers_from_the_books_offline():
    slug = _onboard("Asking Traders")
    _stock(slug)
    resp = client.post(f"/app/{slug}/assistant", data={"question": "Who owes me the most?"})
    assert resp.status_code == 200
    assert 'data-testid="assistant-answer"' in resp.text
    assert "Basavaraj" in resp.text


def test_the_styleguide_is_there_for_operators():
    resp = client.get("/styleguide")
    assert resp.status_code == 200 and 'data-testid="styleguide"' in resp.text


# ======================================================= who may see what

def test_another_account_cannot_open_the_business_or_its_setup():
    slug = _onboard("Private Traders")
    _login(OTHER)
    try:
        site = client.get(f"/app/{slug}")
        assert f"/app/{slug}" not in str(site.url), site.url
        assert "Private Traders" not in site.text
        setup = client.get(f"/studio/{slug}/profile")
        assert str(setup.url).split("?")[0].endswith("/studio"), setup.url
        assert "Private Traders" not in setup.text
        before = store.get_client(slug, ACCOUNT.id).name
        client.post(f"/studio/{slug}/profile", data={"name": "Hijacked"})
        assert store.get_client(slug, ACCOUNT.id).name == before
    finally:
        _login(ACCOUNT)


def test_a_tenant_is_kept_out_of_the_studio():
    _install(OTHER, "tenant")
    _login(OTHER)
    try:
        resp = client.get("/studio")
        assert "/studio" not in str(resp.url), resp.url
    finally:
        _install(OTHER, "operator")
        _login(ACCOUNT)


def test_signed_out_visitors_are_sent_to_login():
    slug = _onboard("Locked Traders")
    anon = TestClient(app_mod.app, follow_redirects=True)
    for path in (f"/app/{slug}", "/studio", f"/studio/{slug}/profile", "/styleguide"):
        resp = anon.get(path)
        assert str(resp.url).endswith("/login"), (path, resp.url)


# ------------------------------------------------------------------ runner

def _cleanup() -> None:
    for slug in set(_SLUGS):
        store.delete_client(slug, ACCOUNT.id)
    for acct in (ACCOUNT, OTHER):
        auth.delete(acct.id)


def main() -> int:
    tests = [(n, f) for n, f in sorted(globals().items())
             if n.startswith("test_") and callable(f)]
    passed = failed = 0
    try:
        for name, fn in tests:
            try:
                fn()
                print(f"ok   {name}")
                passed += 1
            except Exception as exc:                       # noqa: BLE001
                print(f"FAIL {name}: {type(exc).__name__}: {exc}")
                failed += 1
    finally:
        _cleanup()
    print(f"\n{passed} passed, {failed} failed")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
